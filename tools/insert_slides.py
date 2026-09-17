# -*- coding: utf-8 -*-
"""생성기가 만든 장을 사용자 편집본(발표용)에 끼워 넣는다 — 지우지 않고 추가만 한다(R175).

  python3 tools/insert_slides.py [--drop-separators] <생성기 임시 출력.pptx> <사용자 발표용.pptx> "제목:after=앞 장 제목" "제목:before=뒤 장 제목" …

  - 발표용은 09-17 부터 사용자 편집본이 기준이라 생성기 출력으로 덮지 않는다
  - 원본 장의 도형을 deepcopy 로 복사(그림 없는 장만). 배경·노트도 복사. 페이지 번호는 1부터 다시 매긴다
  - 사용자가 지운 것과 같은 종류의 얇은 구분선(격자 괘선이 아닌 높이 0 선)은 복사하지 않는다 — 사용자 덱에 그런 선이 남아 있으면 그대로 복사한다
"""
import copy, os, re, shutil, subprocess, sys, unicodedata
from pptx import Presentation

IN = 914400
ARGS = [a for a in sys.argv[1:] if a != '--drop-separators']
FORCE_DROP = '--drop-separators' in sys.argv      # 사용자가 표·목록 장의 구분선은 지우고 화면 말풍선 장의 것만 남겼을 때
SRC, DST = ARGS[0], ARGS[1]
SPECS = ARGS[2:]

def powerpoint_has(name):
    try:
        out = subprocess.run(['osascript', '-e', 'tell application "Microsoft PowerPoint" to return name of every presentation'], capture_output=True, text=True, timeout=10).stdout
    except Exception:
        return False
    return unicodedata.normalize('NFD', os.path.basename(name)) in unicodedata.normalize('NFD', out)

def title_of(sl):
    for sh in sl.shapes:
        if sh.has_text_frame and abs(sh.top / IN - 0.3) < 0.06 and sh.left / IN < 0.7 and sh.width / IN > 2:
            return sh.text_frame.text.strip()
    return None

def is_hairline(sh):
    return sh.height == 0 and sh.width > 0
def on_grid(sh):
    y, w = sh.top / IN, sh.width / IN
    k = (y - 0.9) / 0.25
    return abs(k - round(k)) < 0.02 and abs(w - 8.8) < 0.05

if powerpoint_has(DST):
    sys.exit('중단: PowerPoint 가 발표용 파일을 열고 있다. 닫고 다시 실행할 것')
TMP = DST + '.tmp'
shutil.copyfile(DST, TMP)
src, dst = Presentation(SRC), Presentation(TMP)
src_by_title = {title_of(s): s for s in src.slides}
off_grid_in_dst = sum(1 for s in dst.slides for sh in s.shapes if is_hairline(sh) and not on_grid(sh))
drop_separators = FORCE_DROP or off_grid_in_dst == 0
print('사용자 덱의 격자 밖 구분선: %d → 구분선 %s' % (off_grid_in_dst, '제외' if drop_separators else '유지'))

lst = dst.slides._sldIdLst
layout = dst.slides[0].slide_layout
for spec in SPECS:
    m = re.match(r'(.+?):(after|before)=(.+)', spec)
    title, mode, anchor = m.group(1), m.group(2), m.group(3)
    s0 = src_by_title.get(title)
    if s0 is None: sys.exit('중단: 원본에 "%s" 장이 없다' % title)
    if any(sh.shape_type == 13 for sh in s0.shapes): sys.exit('중단: "%s" 장에 그림이 있다. 이 도구는 글·도형만 복사한다' % title)
    new = dst.slides.add_slide(layout)
    for ph in list(new.placeholders): ph._element.getparent().remove(ph._element)
    bg = s0._element.cSld.find('{http://schemas.openxmlformats.org/presentationml/2006/main}bg')
    if bg is not None: new._element.cSld.insert(0, copy.deepcopy(bg))
    for sh in s0.shapes:
        if drop_separators and is_hairline(sh) and not on_grid(sh): continue
        new.shapes._spTree.append(copy.deepcopy(sh._element))
    if s0.has_notes_slide and s0.notes_slide.notes_text_frame is not None and s0.notes_slide.notes_text_frame.text.strip():
        ns = new.notes_slide
        if ns.notes_text_frame is None:   # pptxgenjs 덱의 노트 마스터엔 본문 자리표시자가 없다 → 원본 노트 장의 본문 도형을 복사
            for sp in s0.notes_slide.shapes:
                if sp.is_placeholder and sp.has_text_frame: ns.shapes._spTree.append(copy.deepcopy(sp._element))
        if ns.notes_text_frame is not None: ns.notes_text_frame.text = s0.notes_slide.notes_text_frame.text
        else: sys.exit('중단: "%s" 노트를 심을 자리표시자를 만들지 못했다' % title)
    titles = [title_of(s) for s in dst.slides]
    if anchor not in titles: sys.exit('중단: 발표용에 "%s" 장이 없다' % anchor)
    idx = titles.index(anchor) + (1 if mode == 'after' else 0)
    el = lst[-1]; lst.remove(el); lst.insert(idx, el)
    print('  끼움: "%s" → %d 번째 (%s %s)' % (title, idx + 1, mode, anchor))

n = 0
for i, sl in enumerate(dst.slides, 1):
    for sh in sl.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().isdigit() and sh.left / IN > 8.5 and sh.top / IN > 5.0:
            for para in sh.text_frame.paragraphs:
                for r in para.runs: r.text = ''
            sh.text_frame.paragraphs[0].runs[0].text = str(i); n += 1
dst.save(TMP); os.replace(TMP, DST)
print('완료 → %s (%d 장 · 페이지 번호 %d 곳)' % (DST, len(dst.slides), n))
