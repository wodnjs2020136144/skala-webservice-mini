# -*- coding: utf-8 -*-
"""생성된 덱을 검사한다 — python3 tools/deck_check.py [pptx 경로]

  1) 도형이 슬라이드 밖으로 나갔는가
  2) 표·텍스트가 아래로 넘치는가 (행 높이는 PowerPoint 가 정하므로 추정한다)
  3) 자리표시자·철회된 표현이 남았는가
  4) 말풍선에 적은 API 주소가 OAS 에 실재하는가
"""
import io, math, os, re, sys, unicodedata
from pptx import Presentation
try:
    import yaml
except ImportError:
    sys.exit('PyYAML 이 필요하다:  python3 -m pip install pyyaml')

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PPTX = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, 'deliverables', '10반_황재원_너와나의인계고리-개요.pptx')   # 인자로 다른 덱도 검사한다
OAS  = os.path.join(ROOT, 'deliverables', '10반_황재원_너와나의인계고리-API.yml')
IN = 914400
SW, SH = 10.0, 5.625

pr = Presentation(PPTX)
spec = yaml.safe_load(io.open(OAS, encoding='utf-8'))
paths = set(spec['paths'])
fails = []

def wide(c): return unicodedata.east_asian_width(c) in ('W', 'F')
def lines(text, w_in, fs):
    avail = w_in * 72
    n = 0
    for para in text.split('\n'):
        wpt = sum((fs if wide(c) else fs * 0.52) for c in para)
        n += max(1, math.ceil(wpt / avail)) if avail > 0 else 1
    return n

# 1·2) 기하
for i, s in enumerate(pr.slides, 1):
    for sh in s.shapes:
        if sh.left is None: continue
        l, t, w, h = sh.left / IN, sh.top / IN, sh.width / IN, sh.height / IN
        if l < -0.01 or t < -0.01 or l + w > SW + 0.01 or t + h > SH + 0.01:
            fails.append('슬라이드 %d: 도형이 슬라이드를 벗어남 (%s)' % (i, sh.shape_type))
        if sh.has_text_frame and sh.text_frame.text.strip():
            fs = max([r.font.size.pt for pg in sh.text_frame.paragraphs for r in pg.runs if r.font.size] or [12])
            need = lines(sh.text_frame.text, w, fs) * fs * 1.22 / 72
            if t + need > SH - 0.06:
                fails.append('슬라이드 %d: 글이 아래로 넘침 — %s' % (i, sh.text_frame.text[:34].replace('\n', ' ')))
        if sh.has_table:
            cols = [c.width / IN for c in sh.table.columns]
            tot = 0
            for ri, row in enumerate(sh.table.rows):
                mx, fs = 1, 11 if ri == 0 else 10.5
                for ci, cell in enumerate(row.cells):
                    for pg in cell.text_frame.paragraphs:
                        for r in pg.runs:
                            if r.font.size: fs = r.font.size.pt
                    mx = max(mx, lines(cell.text, cols[ci], fs))
                tot += (mx * fs * 1.25 + 10) / 72
            if t + tot > SH - 0.06:
                fails.append('슬라이드 %d: 표가 아래로 넘침 (추정 하단 %.2f)' % (i, t + tot))

# 3) 금지 표현 — 철회 슬라이드에서 인용한 것은 허용한다
QUOTED_OK = {'출근해서야', '역할별 인계', 'care_assignment', '인계 내용이 기록되지 않는다', '전산 기록 하나가'}
BAN = ['[ 채울 것 ]', '11개 엔티티', '23 오퍼레이션', '분기 A/B', '읽기 3단', 'append-only', '서버 계산']
texts = []
for i, s in enumerate(pr.slides, 1):
    buf = []
    for sh in s.shapes:
        if sh.has_text_frame: buf.append(sh.text_frame.text)
        if sh.has_table:
            for row in sh.table.rows:
                for c in row.cells: buf.append(c.text)
    texts.append((i, '\n'.join(buf)))
joined = '\n'.join(t for _, t in texts)
for b in BAN:
    if b in joined:
        where = [i for i, t in texts if b in t]
        fails.append('금지/내부 용어가 남음: "%s" (슬라이드 %s)' % (b, where))

# 4) 말풍선 주소가 OAS 에 있는가
def norm(p):
    p = re.sub(r'\{[^}]*\}', '{x}', p)
    return re.sub(r'\?.*$', '', p)
oas_norm = set(norm(p) for p in paths)
seen, bad = set(), []
for i, t in texts:
    for m in re.finditer(r'(GET|POST|PATCH|PUT|DELETE)\s+(/[A-Za-z0-9_{}/\-…?=]+)', t):
        raw = m.group(2)
        if '…' in raw:            # 줄임 표기는 뒷부분만 대조
            tail = raw.split('…')[-1]
            if not any(op.endswith(norm(tail)) for op in oas_norm):
                bad.append('슬라이드 %d: %s' % (i, m.group(0)))
            continue
        if norm(raw) not in oas_norm:
            bad.append('슬라이드 %d: %s' % (i, m.group(0)))
        seen.add(norm(raw))
for b in bad:
    fails.append('OAS 에 없는 주소를 슬라이드가 적음 — ' + b)

print('슬라이드 %d · 표 %d · 이미지 %d' % (
    len(pr.slides),
    sum(1 for s in pr.slides for sh in s.shapes if sh.has_table),
    sum(1 for s in pr.slides for sh in s.shapes if sh.shape_type == 13)))
if fails:
    print('\n문제 %d 건:' % len(fails))
    for f in fails: print('  -', f)
    sys.exit(1)
print('통과 — 넘침 0 · 금지 표현 0 · 슬라이드의 모든 API 주소가 명세에 있다')
